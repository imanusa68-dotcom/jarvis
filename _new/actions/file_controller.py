# actions/file_controller.py
# File management — create, delete, move, rename, list, find, organize

import hashlib
import shutil
from pathlib import Path
from datetime import datetime
import send2trash


# ── Stage 2.4b: transactional fileops routing ─────────────────
# When the feature flag is ON, mutating actions go through core.fileops so they
# are journaled, backed up (staging) and reversible. Each helper returns None to
# signal "fall back to the legacy implementation below" (flag OFF, out of
# scope, Office doc, or any error) so old behaviour is preserved exactly.

def _get_fileops():
    try:
        from actions.fileops_bridge import get_fileops
        return get_fileops()
    except Exception:
        return None


def _fo_create(full, content, confirmed=False):
    fo = _get_fileops()
    if fo is None:
        return None
    target = _ensure_file_target(Path(full).expanduser(), "new_file.txt")
    if not _is_within_safe_roots(target) or target.suffix.lower() in _OFFICE_EXTS:
        return None
    # Stage 2 decision: creating OVER an existing file must confirm first, so
    # «создай файл» never silently destroys data the user forgot about.
    if target.exists() and target.is_file() and not confirmed:
        try:
            prev = Path(target).read_text(encoding="utf-8", errors="replace")
        except Exception:
            prev = ""
        prev = prev.replace("\r\n", "\n").strip()
        if len(prev) > 80:
            prev = prev[:80].rstrip() + "\u2026"
        shown = f" (сейчас в нём: «{prev}»)" if prev else ""
        return (f"Файл {target.name} уже существует{shown}. Перезаписать его? "
                f"Скажите 'да' для подтверждения.")
    try:
        res = fo.replace_file(target, content or "")
        verb = "Created" if res["created"] else "Overwrote"
        return f"{verb} file: {target.name} (full path: {target})"
    except Exception as e:
        print(f"[fileops] create fell back to legacy: {e}")
        return None


def _fo_write(full, content, append):
    if append:
        return None
    fo = _get_fileops()
    if fo is None:
        return None
    target = _ensure_file_target(Path(full).expanduser(), "new_file.txt")
    if not _is_within_safe_roots(target) or target.suffix.lower() in _OFFICE_EXTS:
        return None
    try:
        res = fo.replace_file(target, content or "")
        verb = "Wrote" if res["created"] else "Overwrote"
        return f"{verb} file: {target.name} (full path: {target})"
    except Exception as e:
        print(f"[fileops] write fell back to legacy: {e}")
        return None


def _fo_move(full, destination):
    fo = _get_fileops()
    if fo is None:
        return None
    src = Path(full).expanduser()
    if not src.exists() or not _is_within_safe_roots(src):
        return None
    dst = _resolve_path(destination)
    if dst.is_dir():
        dst = dst / src.name
    if not _is_within_safe_roots(dst) or dst.exists():
        return None
    try:
        fo.move(src, dst)
        return f"Moved: {src.name} → {dst} (full path: {dst})"
    except Exception as e:
        print(f"[fileops] move fell back to legacy: {e}")
        return None


def _fo_rename(full, new_name):
    fo = _get_fileops()
    if fo is None:
        return None
    src = Path(full).expanduser()
    if not src.exists() or not _is_within_safe_roots(src) or not (new_name or "").strip():
        return None
    safe_name = Path(new_name.strip()).name
    dst = src.with_name(safe_name)
    if not _is_within_safe_roots(dst) or dst.exists():
        return None
    try:
        fo.rename(src, safe_name)
        return f"Renamed: {src.name} → {safe_name} (full path: {dst})"
    except Exception as e:
        print(f"[fileops] rename fell back to legacy: {e}")
        return None


def _fo_copy(full, destination):
    fo = _get_fileops()
    if fo is None:
        return None
    src = Path(full).expanduser()
    if not src.exists() or not src.is_file() or not _is_within_safe_roots(src):
        return None
    if not (destination or "").strip():
        return None
    dst = _resolve_path(destination)
    if dst.is_dir():
        dst = dst / src.name
    if not _is_within_safe_roots(dst) or dst.exists():
        return None
    try:
        fo.copy(src, dst)
        return f"Copied: {src.name} → {dst} (full path: {dst})"
    except Exception as e:
        print(f"[fileops] copy fell back to legacy: {e}")
        return None


def _fo_create_folder(full):
    fo = _get_fileops()
    if fo is None:
        return None
    target = Path(full).expanduser()
    if target.exists() or not _is_within_safe_roots(target):
        return None
    try:
        res = fo.create_folder(target)
        return res["message"]
    except Exception as e:
        print(f"[fileops] create_folder fell back to legacy: {e}")
        return None


def _fo_delete(full):
    fo = _get_fileops()
    if fo is None:
        return None
    target = Path(full).expanduser()
    if not target.exists() or not target.is_file():
        return None
    if not _is_within_safe_roots(target):
        return None
    try:
        res = fo.delete(target)
        return res["message"]
    except Exception as e:
        print(f"[fileops] delete refused/fell back: {e}")
        return None


def _fmt_state(cur):
    """Human phrase for the file's ACTUAL current content after an op."""
    if not cur:
        return ""
    if not cur.get("exists"):
        return " Файла сейчас нет на диске."
    p = cur.get("preview")
    if p is None:
        return " Файл на месте (содержимое нетекстовое)."
    return f" Сейчас в файле: «{p}»."


def _fmt_steps(r):
    fwd = r.get("redo_available", 0)
    back = r.get("undo_available", 0)
    bits = []
    if fwd:
        bits.append(f"вперёд (redo) ещё {fwd}")
    if back:
        bits.append(f"назад (undo) ещё {back}")
    return (" Шаги: " + "; ".join(bits) + ".") if bits else ""


def _scope_target(path, name):
    """Resolve the concrete FILE an undo/redo/history call is about, or None for
    a global (folder-level / unscoped) request. Scoping keeps navigation inside
    the file the user named, so undo can't wander into another file's - or
    another session's - history. Uses the module-level _resolve_path (mirrors the
    dispatcher's nested _full_path) so it works when called outside the closure.
    """
    try:
        base = _resolve_path(path)
    except Exception:
        return None
    n = (name or "").strip()
    if n:
        # Do NOT re-append name when `path` already points at the file (issue 015).
        if base.name == n or (base.exists() and base.is_file()):
            target = base
        else:
            target = base / n
    else:
        target = base
    try:
        if target.is_dir():
            return None
    except Exception:
        pass
    if n or target.suffix:
        return str(target)
    return None


def _fo_undo(scope=None):
    fo = _get_fileops()
    if fo is None:
        return None
    try:
        r = fo.undo_last(scope)
    except Exception as e:
        print(f"[fileops] undo fell back to legacy: {e}")
        return None
    if r is None:
        return None
    if not r["ok"]:
        return f"Не удалось отменить: {r['label']}"
    kind = r.get("kind")
    state = _fmt_state(r.get("current"))
    steps = _fmt_steps(r)
    if kind == "remove-file":
        # Undoing a file's ORIGINAL creation deletes the file. Say so plainly.
        return f"Отменено создание — файл удалён ({r['label']}).{state}{steps}"
    if kind == "move":
        return f"Отменено — перемещение/переименование возвращено ({r['label']}).{state}{steps}"
    return f"Отменено — восстановлено прежнее состояние ({r['label']}).{state}{steps}"


def _fo_redo(scope=None):
    """Re-apply the most recently undone file action (symmetric with _fo_undo)."""
    fo = _get_fileops()
    if fo is None:
        return None
    try:
        r = fo.redo_last(scope)
    except Exception as e:
        print(f"[fileops] redo error: {e}")
        return None
    if r is None:
        return None
    if not r["ok"]:
        return f"Не удалось вернуть вперёд: {r['label']}"
    kind = r.get("kind")
    state = _fmt_state(r.get("current"))
    steps = _fmt_steps(r)
    if kind == "move":
        return f"Возвращено вперёд — перемещение/переименование применено заново ({r['label']}).{state}{steps}"
    return f"Возвращено вперёд — изменение применено заново ({r['label']}).{state}{steps}"


def _fo_history(scope=None):
    """Render a content-aware version timeline for the model and the user."""
    fo = _get_fileops()
    if fo is None:
        return None
    try:
        h = fo.history(path=scope)
    except Exception as e:
        print(f"[fileops] history error: {e}")
        return None
    cur = h.get("current")
    forward = h.get("forward") or []
    back = h.get("back") or []
    if cur is None and not forward and not back:
        return None
    lines = []
    if cur:
        if not cur.get("exists"):
            lines.append("Сейчас: файла нет на диске.")
        else:
            pv = cur.get("preview")
            lines.append(f"Сейчас в файле: «{pv}»." if pv is not None
                         else "Сейчас: файл на месте.")
    def _desc(e):
        if e.get("preview") is not None:
            return f"станет «{e['preview']}»"
        return e.get("note") or "изменение"
    if forward:
        lines.append("Вперёд (redo), шаг за шагом:")
        for i, e in enumerate(forward, 1):
            lines.append(f"  {i}. {e.get('label')} → {_desc(e)}")
    if back:
        lines.append("Назад (undo), шаг за шагом:")
        for i, e in enumerate(back, 1):
            lines.append(f"  {i}. {e.get('label')} → {_desc(e)}")
    return "\n".join(lines)



# ── Undo backups (issue 013) ─────────────────────────────────────────────────
# Before OVERWRITING an existing file we stash its previous bytes here, so
# "верни как было" can restore it. Lives outside the user's folders (its own
# hidden dir) so it never clutters Desktop/Documents. Single-level (last
# overwrite) — enough for the "undo my last change" case.

def _backup_dir() -> Path:
    """Where previous file versions are stashed.

    Шаг 35.1: шов был обещан в этой строке документации, но не сделан —
    путь считался от Path.home() и в тестах вёл в настоящий дом. Теперь
    дом один на всех: core.safe_json.state_dir().
    """
    from core.safe_json import state_dir
    return state_dir() / "backups"


def _backup_key(target: Path) -> str:
    """Stable per-path backup file name (hash keeps it flat and filesystem-safe)."""
    h = hashlib.sha1(str(target.resolve()).encode("utf-8", "ignore")).hexdigest()[:16]
    return f"{h}{target.suffix}"


def _save_backup(target: Path) -> None:
    """Copy the CURRENT bytes of `target` to the backup store. Best-effort:
    a backup failure must never block the write itself."""
    try:
        bdir = _backup_dir()
        bdir.mkdir(parents=True, exist_ok=True)
        shutil.copy2(str(target), str(bdir / _backup_key(target)))
    except Exception as e:
        print(f"[FileController] backup skipped: {e}")


def restore_file(path: str) -> str:
    """Restore a file from its last backup — the 'верни как было' action."""
    target = Path(path).expanduser()
    if not _is_within_safe_roots(target):
        return _refuse_outside_safe_roots(target)
    backup = _backup_dir() / _backup_key(target)
    if not backup.exists():
        return f"Нет сохранённой версии для отмены: {target.name}"
    try:
        shutil.copy2(str(backup), str(target))
        print(f"[FileController] Restored from backup: {target}")
        return f"Восстановлено прежнее содержимое: {target.name} (full path: {target})"
    except Exception as e:
        return f"Не удалось восстановить: {e}"

# Personal folders resolved via the SAME OneDrive-aware source the awareness layer
# uses (watcher/resolver/open_path), so files land where Explorer shows them and
# not in a naive home-relative folder the user never opens (issue 006).
_KNOWN_FOLDER_NAMES = ("desktop", "downloads", "documents", "pictures", "music", "videos")


def _known_folder(name: str) -> Path | None:
    """Real path of a personal folder from core.awareness._known_folders, or None."""
    try:
        from core.awareness import _known_folders
        return _known_folders.folder(name)
    except Exception:
        return None


def _get_desktop() -> Path:
    """Returns desktop path — works on Windows, Mac, Linux."""
    return _known_folder("desktop") or (Path.home() / "Desktop")


def _get_downloads() -> Path:
    return _known_folder("downloads") or (Path.home() / "Downloads")


def _resolve_path(raw: str) -> Path:
    """
    Resolves a path from user input.
    Supports shortcuts: 'desktop', 'downloads', 'documents', 'pictures', 'music',
    'videos', 'home' — the personal folders resolve to their REAL (redirect-aware)
    location so file ops match what the user sees in Explorer.
    """
    raw = (raw or "").strip()
    lower = raw.lower()
    if lower == "home":
        return Path.home()
    if lower in _KNOWN_FOLDER_NAMES:
        return _known_folder(lower) or (Path.home() / lower.capitalize())

    # Shortcut + subpath, e.g. "documents/Мои сочинения", "desktop\\proj\\sub" →
    # resolve the subpath INSIDE the real known folder (not a cwd-relative path).
    norm = raw.replace("\\", "/")
    head, sep, tail = norm.partition("/")
    if sep and head.lower() in _KNOWN_FOLDER_NAMES:
        base = _known_folder(head.lower()) or (Path.home() / head.capitalize())
        parts = [s for s in tail.split("/") if s]
        return base.joinpath(*parts) if parts else base

    p = Path(raw).expanduser()
    if not p.is_absolute():
        # A bare/relative name must NEVER resolve against the app's working
        # directory — that silently buried files inside the project tree
        # (issue 018). Anchor it to the Desktop, where users keep their folders.
        desktop = _known_folder("desktop") or (Path.home() / "Desktop")
        return desktop.joinpath(*[s for s in norm.split("/") if s])
    return p


def _format_size(bytes_size: int) -> str:
    """Converts bytes to human readable format."""
    for unit in ["B", "KB", "MB", "GB", "TB"]:
        if bytes_size < 1024:
            return f"{bytes_size:.1f} {unit}"
        bytes_size /= 1024
    return f"{bytes_size:.1f} TB"


# ─────────────────────────────────────────────────────────────────────────────
# Stage 3 — path boundary guard.
# Write/create operations are confined to a small set of user folders. System
# locations (Windows, Program Files, drive roots) and any path escaping the
# allowed roots (via "..", symlinks, etc.) are refused. Read-only operations are
# intentionally NOT restricted by this — only mutating operations are.
# ─────────────────────────────────────────────────────────────────────────────

def _safe_roots() -> list[Path]:
    """
    User folders where mutating file operations are permitted — resolved via the
    same known-folder source as _resolve_path, so the boundary matches where files
    are actually created/moved (OneDrive-aware). Deduplicated.
    """
    roots: list[Path] = []
    seen: set[str] = set()
    for name in _KNOWN_FOLDER_NAMES:
        p = _known_folder(name) or (Path.home() / name.capitalize())
        key = str(p).lower()
        if key not in seen:
            seen.add(key)
            roots.append(p)
    return roots


def _is_within_safe_roots(target: Path) -> bool:
    """
    True if `target` resolves to a location inside one of the safe roots.
    Resolves the path first so that '..' traversal and symlinks cannot escape.
    The target itself need not exist (we resolve its nearest existing ancestor).
    """
    try:
        probe = target
        # For a not-yet-existing file, walk up to the nearest existing ancestor
        # so a missing leaf does not defeat the check.
        while not probe.exists() and probe != probe.parent:
            probe = probe.parent
        resolved = probe.resolve()
    except Exception:
        return False

    for root in _safe_roots():
        try:
            root_resolved = root.resolve()
        except Exception:
            continue
        if resolved == root_resolved or root_resolved in resolved.parents:
            return True
    return False


def _refuse_outside_safe_roots(target: Path) -> str:
    """Standard refusal message for a mutating op outside the safe roots."""
    allowed = ", ".join(r.name for r in _safe_roots())
    return (
        f"SECURITY: For safety, files can only be created or written inside your "
        f"personal folders ({allowed}). Refused path: {target}"
    )


def _ensure_file_target(target: Path, default_name: str) -> Path:
    """
    Make `target` a FILE path, never a folder.

    The model sometimes omits the file name and passes only a folder (e.g.
    'desktop'). Writing to a directory would raise a confusing PermissionError,
    so instead of failing we keep Jarvis capable: drop a sensibly-named file
    inside that folder. Returns the resolved file path to use.
    """
    # Existing directory, or a shortcut that resolves to one of our roots →
    # treat as "folder given, name missing" and use the default file name.
    if target.is_dir():
        return target / default_name
    return target


# Office formats handled by a real library, not plain-text write.
_OFFICE_EXTS = {".docx", ".xlsx"}


def _write_office_document(target: Path, content: str) -> str:
    """
    Create a REAL Office document (.docx / .xlsx) — not a renamed text file.
    A .docx/.xlsx is a structured ZIP, so plain write_text() would corrupt it.
    Falls back to a clear message if the optional library is unavailable.
    """
    ext = target.suffix.lower()
    try:
        if ext == ".docx":
            from docx import Document  # python-docx
            doc = Document()
            for line in (content or "").split("\n"):
                doc.add_paragraph(line)
            doc.save(str(target))
        elif ext == ".xlsx":
            from openpyxl import Workbook
            wb = Workbook()
            ws = wb.active
            # Each line → a row; tabs/commas split into columns.
            for r, line in enumerate((content or "").split("\n"), start=1):
                cells = line.split("\t") if "\t" in line else line.split(",")
                for c, val in enumerate(cells, start=1):
                    ws.cell(row=r, column=c, value=val)
            wb.save(str(target))
        else:
            return f"Unsupported office format: {ext}"
        print(f"[FileController] Created office document: {target}")
        return f"Created document: {target.name} (full path: {target})"
    except ImportError:
        lib = "python-docx" if ext == ".docx" else "openpyxl"
        return (
            f"Cannot create {ext} files — the '{lib}' library is not installed. "
            f"Install it, or I can create a plain .txt instead."
        )
    except Exception as e:
        return f"Could not create document: {e}"


def list_files(path: str = "desktop", show_hidden: bool = False) -> str:
    """Lists files and folders in a directory."""
    try:
        target = _resolve_path(path)
        if not target.exists():
            return f"Path not found: {target}"
        if not target.is_dir():
            return f"Not a directory: {target}"

        items = []
        for item in sorted(target.iterdir()):
            if not show_hidden and item.name.startswith("."):
                continue
            if item.is_dir():
                items.append(f"📁 {item.name}/")
            else:
                size = _format_size(item.stat().st_size)
                items.append(f"📄 {item.name} ({size})")

        if not items:
            return f"Directory is empty: {target}"

        return f"Contents of {target.name}/ ({len(items)} items):\n" + "\n".join(items)

    except PermissionError:
        return f"Permission denied: {path}"
    except Exception as e:
        return f"Error listing files: {e}"


def create_file(path: str, content: str = "") -> str:
    """
    Create a new file (optionally with content), inside the safe user folders.
    Refuses system locations and paths that escape the allowed roots.
    """
    try:
        target = Path(path).expanduser()
        if not _is_within_safe_roots(target):
            print(f"[FileController] REFUSED create outside safe roots: {target}")
            return _refuse_outside_safe_roots(target)

        # No file name given (folder only) → drop a sensibly-named file inside.
        target = _ensure_file_target(target, "new_file.txt")

        # Real Office documents need a proper library, not write_text.
        if target.suffix.lower() in _OFFICE_EXTS:
            target.parent.mkdir(parents=True, exist_ok=True)
            return _write_office_document(target, content)

        existed = target.exists()
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_text(content or "", encoding="utf-8")

        verb = "Overwrote" if existed else "Created"
        print(f"[FileController] {verb} file: {target}")
        return f"{verb} file: {target.name} (full path: {target})"
    except PermissionError:
        return f"Permission denied: {path}"
    except Exception as e:
        return f"Could not create file: {e}"


def create_folder(path: str) -> str:
    """
    Create a folder inside the safe user folders.
    Refuses system locations and paths that escape the allowed roots.
    """
    try:
        target = Path(path).expanduser()
        if not _is_within_safe_roots(target):
            print(f"[FileController] REFUSED folder outside safe roots: {target}")
            return _refuse_outside_safe_roots(target)

        # Only a parent folder was given (no new-folder name) → use a default
        # name inside it instead of reporting "already exists".
        if target.is_dir():
            target = target / "new_folder"

        if target.exists():
            return f"Folder already exists: {target.name}"
        target.mkdir(parents=True, exist_ok=True)
        print(f"[FileController] Created folder: {target}")
        return f"Created folder: {target.name} (full path: {target})"
    except PermissionError:
        return f"Permission denied: {path}"
    except Exception as e:
        return f"Could not create folder: {e}"


def delete_file(path: str, confirm: bool = True) -> str:
    """
    SECURED: File deletion is disabled for safety.
    """
    print(f"[FileController] BLOCKED: Delete operation disabled for: {path}")
    return (
        "SECURITY: File deletion is disabled for safety. "
        f"Attempted to delete: {Path(path).name}"
    )


def move_file(source: str, destination: str) -> str:
    """
    Move a file/folder. Both source and destination must be inside the safe
    user folders. Refuses to overwrite an existing destination.
    """
    try:
        src = Path(source).expanduser()
        if not src.exists():
            return f"Source not found: {source}"
        if not _is_within_safe_roots(src):
            return _refuse_outside_safe_roots(src)

        dst = _resolve_path(destination)
        # If destination is a folder, keep the original file name inside it.
        if dst.is_dir():
            dst = dst / src.name
        if not _is_within_safe_roots(dst):
            return _refuse_outside_safe_roots(dst)
        if dst.exists():
            return f"Destination already exists: {dst.name} (not overwriting)"

        dst.parent.mkdir(parents=True, exist_ok=True)
        shutil.move(str(src), str(dst))
        print(f"[FileController] Moved: {src} -> {dst}")
        return f"Moved: {src.name} → {dst} (full path: {dst})"
    except PermissionError:
        return f"Permission denied: {source}"
    except Exception as e:
        return f"Could not move: {e}"


def copy_file(source: str, destination: str) -> str:
    """Copies a file or folder."""
    try:
        src = Path(source).expanduser()
        dst = _resolve_path(destination)

        if not src.exists():
            return f"Source not found: {source}"

        if dst.is_dir():
            dst = dst / src.name

        dst.parent.mkdir(parents=True, exist_ok=True)

        if src.is_dir():
            shutil.copytree(str(src), str(dst))
        else:
            shutil.copy2(str(src), str(dst))

        return f"Copied: {src.name} → {dst.parent.name}/"

    except Exception as e:
        return f"Could not copy: {e}"


def rename_file(path: str, new_name: str) -> str:
    """
    Rename a file/folder in place. The target must be inside the safe user
    folders. new_name is a bare name (no path) — renaming never moves across
    folders. Refuses to overwrite an existing sibling.
    """
    try:
        src = Path(path).expanduser()
        if not src.exists():
            return f"Not found: {path}"
        if not _is_within_safe_roots(src):
            return _refuse_outside_safe_roots(src)
        if not new_name or not new_name.strip():
            return "Please provide a new name."

        # Keep it a bare name — strip any path components for safety.
        safe_name = Path(new_name.strip()).name
        dst = src.with_name(safe_name)
        if not _is_within_safe_roots(dst):
            return _refuse_outside_safe_roots(dst)
        if dst.exists():
            return f"A file named '{safe_name}' already exists (not overwriting)."

        src.rename(dst)
        print(f"[FileController] Renamed: {src} -> {dst}")
        return f"Renamed: {src.name} → {safe_name} (full path: {dst})"
    except PermissionError:
        return f"Permission denied: {path}"
    except Exception as e:
        return f"Could not rename: {e}"


def read_file(path: str, max_chars: int = 3000) -> str:
    """Reads and returns the content of a text file."""
    try:
        target = Path(path).expanduser()
        if not target.exists():
            folder = str(target.parent) if target.parent else "desktop"
            return f"File not found: {path}" + _suggest_files(target.name, folder)
        if not target.is_file():
            return f"Not a file: {path}"

        content = target.read_text(encoding="utf-8", errors="ignore")
        if len(content) > max_chars:
            content = content[:max_chars] + f"\n\n... (truncated, {len(content)} total chars)"
        return content

    except Exception as e:
        return f"Could not read file: {e}"


def write_file(path: str, content: str, append: bool = False) -> str:
    """
    Write (or append) text to a file inside the safe user folders.
    Refuses system locations and paths that escape the allowed roots.
    """
    try:
        target = Path(path).expanduser()
        if not _is_within_safe_roots(target):
            print(f"[FileController] REFUSED write outside safe roots: {target}")
            return _refuse_outside_safe_roots(target)

        # No file name given (folder only) → write to a sensibly-named file.
        target = _ensure_file_target(target, "new_file.txt")

        # Back up the previous bytes before a true overwrite, so "верни как было"
        # can undo it. Only on overwrite — create/append never lose data.
        if target.exists() and not append:
            try:
                _save_backup(target)
            except Exception as e:
                print(f"[FileController] backup skipped: {e}")

        # Real Office documents need a proper library, not text append/write.
        if target.suffix.lower() in _OFFICE_EXTS:
            target.parent.mkdir(parents=True, exist_ok=True)
            return _write_office_document(target, content)

        existed = target.exists()
        target.parent.mkdir(parents=True, exist_ok=True)
        mode = "a" if append else "w"
        with open(target, mode, encoding="utf-8") as f:
            f.write(content or "")

        if append:
            verb = "Appended to"
        else:
            verb = "Overwrote" if existed else "Wrote"
        print(f"[FileController] {verb} file: {target}")
        return f"{verb} file: {target.name} (full path: {target})"
    except PermissionError:
        return f"Permission denied: {path}"
    except Exception as e:
        return f"Could not write file: {e}"


# Совпадение по содержимому — файл СОДЕРЖИТ этот текст, имя может отличаться
MATCHED_BY_CONTENT_NOTE = (
    "(совпадение по содержимому — файл СОДЕРЖИТ этот текст, "
    "имя может отличаться)"
)


def _suggest_files(query: str, folder: str = "desktop") -> str:
    """'Возможно, вы имели в виду' hint when a named file is not found:
    closest matches by NAME and by CONTENT, so a miss is self-correcting rather
    than a flat 'file does not exist'."""
    q = (query or "").strip()
    if not q:
        return ""
    try:
        from core.awareness import _search
    except Exception:
        return ""
    scope = None
    try:
        rp = _resolve_path(folder)
        if rp.exists():
            scope = rp
    except Exception:
        pass
    seen: set[str] = set()
    cands: list[tuple] = []
    try:
        for r in (_search.search(name=q, scope=scope, limit=5) or []):
            p = Path(r["path"]); k = str(p).lower()
            if k not in seen:
                seen.add(k); cands.append((p, "имени"))
        for r in (_search.search_content(q, scope=scope, limit=5) or []):
            p = Path(r["path"]); k = str(p).lower()
            if k not in seen:
                seen.add(k); cands.append((p, "содержимому"))
    except Exception:
        return ""
    if not cands:
        return ""
    lines = [f"  • {p.name} ({p.parent}) — совпало по {why}" for p, why in cands[:5]]
    return "\nВозможно, вы имели в виду:\n" + "\n".join(lines)


def find_files(name: str = "", extension: str = "", path: str = "home",
               max_results: int = 20, content: str = "") -> str:
    """
    Searches for files by NAME, EXTENSION, or text CONTENT.

    Content matters because users refer to a file by what is INSIDE it
    ('файл привет' = the file containing привет), not its name. If a
    name search finds nothing, we retry by content, so a misremembered name
    still resolves to the real file instead of a dead end.
    """
    try:
        from core.awareness import _search

        # A specific folder scopes the search; the broad ("home"/"pc"/…) case lets
        # the fast index (Everything) cover the whole disk and the fallback cover
        # the user's personal folders. Either way — never an unbounded rglob.
        scope = None
        broad = ("home", "", "pc", "computer", "everywhere", "all",
                 "весь пк", "компьютер", "везде")
        if path and path.strip().lower() not in broad:
            rp = _resolve_path(path)
            if rp.exists():
                scope = rp

        matched_by = "name"
        if content.strip():
            results = _search.search_content(content, scope=scope, limit=max_results)
            matched_by = "content"
        else:
            results = _search.search(name=name, ext=extension, scope=scope,
                                     limit=max_results)
            # Misremembered name? Look at what is INSIDE the files.
            if not results and name.strip() and not extension.strip():
                alt = _search.search_content(name, scope=scope, limit=max_results)
                if alt:
                    results, matched_by = alt, "content"

        if not results:
            query = content or name or extension or "files"
            hint = "" if _search.backend_name() == "everything" else (
                " (искал в ваших основных папках; для поиска по всему ПК "
                "установите Everything)"
            )
            return f"No {query} found{hint}."

        note = (" " + MATCHED_BY_CONTENT_NOTE) if matched_by == "content" else ""
        lines = [f"Found {len(results)} file(s){note}:"]
        for r in results:
            p = Path(r["path"])
            lines.append(f"📄 {p.name} ({_format_size(r['size'])}) — {p.parent}")
        return "\n".join(lines)

    except Exception as e:
        return f"Search error: {e}"


def get_largest_files(path: str = "home", count: int = 10) -> str:
    """Returns the largest files in a directory."""
    try:
        from core.awareness import _search

        scope = None
        broad = ("home", "", "pc", "computer", "everywhere", "all")
        if path and path.strip().lower() not in broad:
            rp = _resolve_path(path)
            if rp.exists():
                scope = rp

        top = _search.largest(scope=scope, count=count)
        if not top:
            return "No files found."

        lines = [f"Top {len(top)} largest files:"]
        for r in top:
            p = Path(r["path"])
            lines.append(f"  {_format_size(r['size']):>10}  {p.name}  ({p.parent})")
        return "\n".join(lines)

    except Exception as e:
        return f"Error: {e}"


def get_disk_usage(path: str = "home") -> str:
    """Returns disk usage information."""
    try:
        target = _resolve_path(path)
        usage  = shutil.disk_usage(target)
        total  = _format_size(usage.total)
        used   = _format_size(usage.used)
        free   = _format_size(usage.free)
        pct    = usage.used / usage.total * 100

        return (
            f"Disk usage for {target}:\n"
            f"  Total : {total}\n"
            f"  Used  : {used} ({pct:.1f}%)\n"
            f"  Free  : {free}"
        )
    except Exception as e:
        return f"Could not get disk usage: {e}"


def organize_desktop() -> str:
    """
    SECURED: Desktop organization (file moving) is disabled for safety.
    """
    print("[FileController] BLOCKED: Organize desktop operation disabled")
    return "SECURITY: Desktop organization is disabled for safety (involves moving files)."


def get_file_info(path: str) -> str:
    """Returns detailed information about a file."""
    try:
        target = Path(path).expanduser()
        if not target.exists():
            return f"Not found: {path}"

        stat = target.stat()
        info = {
            "Name":     target.name,
            "Type":     "Folder" if target.is_dir() else "File",
            "Size":     _format_size(stat.st_size),
            "Location": str(target.parent),
            "Created":  datetime.fromtimestamp(stat.st_ctime).strftime("%Y-%m-%d %H:%M"),
            "Modified": datetime.fromtimestamp(stat.st_mtime).strftime("%Y-%m-%d %H:%M"),
            "Extension": target.suffix or "None",
        }

        return "\n".join(f"  {k}: {v}" for k, v in info.items())

    except Exception as e:
        return f"Could not get file info: {e}"

def file_controller(
    parameters: dict,
    response=None,
    player=None,
    session_memory=None
) -> str:
    action  = (parameters or {}).get("action", "").lower().strip()
    path    = (parameters or {}).get("path", "desktop")
    name    = (parameters or {}).get("name", "")
    content = (parameters or {}).get("content", "")

    def _full_path(p: str, n: str) -> str:
        base = _resolve_path(p)
        if n:
            # Do NOT re-append the name when `path` already points at the file —
            # the model sometimes passes the full path in `path` AND repeats the
            # name in `name`, which used to produce '.../t.txt/t.txt' and crash
            # writes with WinError 183 (issue 015).
            if base.name == n or (base.exists() and base.is_file()):
                return str(base)
            return str(base / n)
        return str(base)

    result = "Unknown action."

    try:
        if action == "list":
            result = list_files(path)

        elif action == "create_file":
            full = _full_path(path, name)
            result = _fo_create(full, content, parameters.get("confirmed", False)) or create_file(full, content=content)

        elif action == "create_folder":
            full = _full_path(path, name)
            result = _fo_create_folder(full) or create_folder(full)

        elif action == "delete":
            full = _full_path(path, name)
            result = _fo_delete(full) or delete_file(full)

        elif action == "move":
            full = _full_path(path, name)
            result = _fo_move(full, parameters.get("destination", "")) or move_file(full, parameters.get("destination", ""))

        elif action == "copy":
            full = _full_path(path, name)
            result = (_fo_copy(full, parameters.get("destination", ""))
                      or copy_file(full, parameters.get("destination", "")))

        elif action == "rename":
            full = _full_path(path, name)
            result = _fo_rename(full, parameters.get("new_name", "")) or rename_file(full, parameters.get("new_name", ""))

        elif action == "read":
            full = _full_path(path, name)
            result = read_file(full)

        elif action in ("undo", "restore"):
            full = _full_path(path, name)
            # Trust the fileops journal (Fix A keeps its safe roots in sync, so
            # writes and undo take the SAME path). Fall back to the legacy
            # single-slot backup only when fileops undo did nothing (None) or
            # could not actually restore (e.g. staged copy evicted by quota) —
            # never let a stale legacy copy shadow a newer fileops version.
            scope = _scope_target(path, name)
            result = _fo_undo(scope)
            if result is None or result.startswith("Не удалось"):
                result = restore_file(full)

        elif action == "redo":
            # Re-apply the last undone action. No legacy fallback exists — the
            # redo stack lives only in the fileops journal.
            scope = _scope_target(path, name)
            result = _fo_redo(scope)
            if result is None:
                result = "Нечего вернуть вперёд — стек повторения пуст."

        elif action == "history":
            # Content-aware timeline so the model can navigate to a NAMED
            # version instead of blind one-step hops.
            scope = _scope_target(path, name)
            result = _fo_history(scope)
            if result is None:
                result = "История изменений пуста."

        elif action == "write":
            full = _full_path(path, name)
            result = _fo_write(full, content, parameters.get("append", False))
            if result is None:
                result = write_file(
                    full,
                    content=content,
                    append=parameters.get("append", False)
                )

        elif action == "find":
            result = find_files(
                name=name or parameters.get("name", ""),
                extension=parameters.get("extension", ""),
                path=path,
                max_results=parameters.get("max_results", 20),
                content=parameters.get("content", ""),
            )

        elif action == "largest":
            result = get_largest_files(
                path=path,
                count=parameters.get("count", 10)
            )

        elif action == "disk_usage":
            result = get_disk_usage(path)

        elif action == "organize_desktop":
            result = organize_desktop()

        elif action == "info":
            full = _full_path(path, name)
            result = get_file_info(full)

        else:
            result = f"Unknown action: '{action}'"

    except Exception as e:
        result = f"File controller error: {e}"

    if player:
        player.write_log(f"[file] {result[:60]}")

    return result
